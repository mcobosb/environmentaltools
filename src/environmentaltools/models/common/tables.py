"""Tables shared by the model workflows: tables of changes (one row per case or variant), Excel output, fingerprints.

A *table of changes* has one row per case (or variant) and one column per changed
parameter, written as a dotted path of the model configuration (e.g.
``hydrodynamics.courant_number`` or ``physics.wcap``). An empty cell means "as in the
base configuration".
"""

import hashlib
import math
from pathlib import Path

import pandas as pd
import yaml

#: Columns of a table of changes that are free text (ignored), besides those starting with ``_``.
NOTE_COLUMNS = {"description", "notes", "note", "comment", "comments"}
CHANGED_FILL = "F8CBAD"   # orange: value different from the base


def parse_cell(v):
    """Value of a cell: numbers as numbers (integers without decimals), dates as ISO text, other text parsed as
    YAML (so ``true``, ``null``-free lists such as ``[[0, 10], [5000, 20]]`` work), empty → ``None``."""
    if v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and not v.strip()):
        return None
    if hasattr(v, "isoformat"):
        return v.isoformat()
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return int(v) if float(v).is_integer() else float(v)
    if hasattr(v, "item"):                                       # numpy scalars
        return parse_cell(v.item())
    text = str(v).strip()
    try:
        parsed = yaml.safe_load(text)
    except yaml.YAMLError:
        return text
    return parsed if isinstance(parsed, (int, float, bool, list, dict)) else text


def set_nested(d, keys, value):
    """``d[k1][k2]…[kn] = value``, creating the intermediate dicts."""
    for k in keys[:-1]:
        d = d.setdefault(k, {})
    d[keys[-1]] = value


def read_table(file_name, sheet_name=None, default_sheet=None):
    """Read an Excel (``sheet_name``, else ``default_sheet`` if it exists, else the first sheet) or csv table."""
    f = Path(file_name)
    if f.suffix.lower() in (".xlsx", ".xls"):
        sheets = pd.ExcelFile(f).sheet_names
        sheet = sheet_name or (default_sheet if default_sheet in sheets else sheets[0])
        df = pd.read_excel(f, sheet_name=sheet)
    else:
        df = pd.read_csv(f)
    df.columns = [str(c).strip() for c in df.columns]
    return df


def read_changes_table(file_name, key="case", sheet_name=None, default_sheet=None, column=None):
    """Read a table of changes: ``{name: nested dict of changes}``.

    Args:
        file_name (str | Path): Excel or csv file.
        key (str): Column with the name of each row (case, variant…).
        sheet_name (str | None): Sheet to read (default: ``default_sheet`` or the first one).
        default_sheet (str | None): Sheet read when ``sheet_name`` is not given and it exists.
        column (callable | None): ``column(name, value, changes, file)`` to handle a column in a
            model-specific way; returns ``True`` if it did. Other columns are dotted paths.

    Returns:
        dict: One entry per row, in table order; empty cells are left out.
    """
    f = Path(file_name)
    df = read_table(f, sheet_name, default_sheet)
    if key not in df:
        raise ValueError(f"{f.name}: the table needs a '{key}' column")
    out = {}
    for _, row in df.iterrows():
        name = parse_cell(row[key])
        if name is None:
            continue
        name = str(name)
        if name in out:
            raise ValueError(f"{f.name}: '{name}' appears twice")
        changes = {}
        for col, v in row.items():
            value = parse_cell(v)
            if col == key or value is None or col.lower() in NOTE_COLUMNS or col.startswith("_"):
                continue
            if column is None or not column(col, value, changes, f):
                set_nested(changes, col.split("."), value)
        out[name] = changes
    return out


def flatten(d, prefix=""):
    """Nested dict → ``{"a.b.c": value}`` (``None`` as empty text)."""
    out = {}
    for k, v in (d or {}).items():
        key = f"{prefix}{k}"
        if isinstance(v, dict) and v:
            out.update(flatten(v, key + "."))
        else:
            out[key] = "" if v is None else v
    return out


def changes_to_table(changes, key="case", descriptions=None):
    """``{name: nested dict of changes}`` → table of changes (one row per name, dotted columns)."""
    import json

    rows = []
    for name, c in changes.items():
        row = {key: name, "description": (descriptions or {}).get(name, "")}
        row.update({k: json.dumps(v) if isinstance(v, (list, dict)) else v for k, v in flatten(c).items()})
        rows.append(row)
    return pd.DataFrame(rows, columns=None if rows else [key, "description"])


def write_tables(path, sheets, legend=None, highlight=None):
    """Write DataFrames as sheets of one Excel file: Arial, bold frozen header, filters, fitted widths and a
    ``legend`` sheet.

    Args:
        path (str | Path): Output file.
        sheets (dict): ``{sheet name: DataFrame}`` (the index is not written).
        legend (dict | None): ``{column or item: description}``, written in the ``legend`` sheet.
        highlight (dict | None): ``{sheet name: (reference column, first compared column)}``: cells of
            the compared columns different from the reference column are filled in orange.

    Returns:
        Path: The written file.
    """
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
        if legend:
            pd.DataFrame({"item": list(legend), "description": list(legend.values())}).to_excel(
                xw, sheet_name="legend", index=False)
        for ws in xw.book.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    c.font = Font(name="Arial", bold=c.row == 1)
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for i, col in enumerate(ws.iter_cols(), start=1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        if legend:
            ws = xw.book["legend"]
            for c in ws["B"]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions["B"].width = 100
        fill = PatternFill("solid", start_color=CHANGED_FILL)
        for name, (ref_col, first_col) in (highlight or {}).items():
            df = sheets[name]
            if ref_col not in df or first_col not in df:
                continue
            r, f = df.columns.get_loc(ref_col), df.columns.get_loc(first_col)
            ws = xw.book[name]
            for i, row in enumerate(df.itertuples(index=False), start=2):
                for j in range(f, len(df.columns)):
                    if str(row[j]) != str(row[r]):
                        ws.cell(i, j + 1).fill = fill
    return path


def folder_signature(folder, skip=()):
    """SHA-1 of the names and contents of the files of a folder (to know whether it has changed)."""
    h = hashlib.sha1()
    for f in sorted(p for p in Path(folder).iterdir() if p.is_file() and p.name not in skip
                    and not p.name.startswith(".")):
        h.update(f.name.encode())
        h.update(f.read_bytes())
    return h.hexdigest()


def is_up_to_date(output, sources):
    """True if ``output`` exists and is newer than every existing file of ``sources``."""
    output = Path(output)
    if not output.exists():
        return False
    t = output.stat().st_mtime
    return all(t > Path(s).stat().st_mtime for s in sources if s and Path(s).exists())
